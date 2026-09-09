from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

SOURCE_XLSX = Path(r"C:\Users\kjahn\Desktop\Söll Gerüstbau Gefährdungsbeurteilung.xlsx")
OUTPUT_XML = Path(r"C:\Users\kjahn\IdeaProjects\GBUXML\src\main\resources\sample-soll-geruestbau-gbuxml.xml")

IGNORE_SHEETS = {
    "Einleitung",
    "Übersichtsblatt",
    "Anhang",
    "Vorlage",
    "Gefährdungskatalog",
    "Mitarbeiterunterweisungen",
    "eGBU",
    "Tabelle1",
    "Tabelle2",
}


def normalise_sheet_name(raw: str) -> str:
    return (raw or "").replace("_", " ").strip()


def parse_date(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    if not text:
        return ""
    text = text.split(",", 1)[0].strip()
    if not text:
        return ""
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", text)
    if m:
        day, month, year = m.groups()
        year = int(year)
        if year < 100:
            year += 2000 if year < 50 else 1900
        try:
            return datetime(year, int(month), int(day)).date().isoformat()
        except ValueError:
            return ""
    return ""


def clean_text(value):
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ").strip()
    while "  " in text:
        text = text.replace("  ", " ")
    return text


def extract_action_needed(value):
    if value is None:
        return False
    text = str(value).strip().lower()
    return text in {"p", "ja", "yes", "x", "true", "1"}


def extract_responsible(value):
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    if "," in text:
        return text.split(",", 1)[1].strip()
    return text


def get_trade_name():
    return "Söll Gerüstbau"


def main():
    wb = load_workbook(str(SOURCE_XLSX), data_only=True)

    root = ET.Element("GBUXML")

    general = ET.SubElement(root, "GeneralData")
    for key, value in [
        ("Firmenname", "Söll Gerüstbau"),
        ("StrasseHausnummer", "Musterstraße 12"),
        ("PLZ", "12345"),
        ("Ort", "Musterstadt"),
        ("ErstellerGefaehrdungsbeurteilung", "Beispielperson"),
        ("DatumDerErstellung", "2025-09-01"),
        ("Gueltigkeitsbereich", "Gefährdungsbeurteilung aus Excel-Vorlage"),
        ("SiFa", "Sicherheitsbeauftragte"),
        ("Betriebsarzt", "Dr. Mustermann"),
    ]:
        child = ET.SubElement(general, key)
        child.text = value

    trades = ET.SubElement(root, "Trades")
    trade = ET.SubElement(trades, "Trade")

    ET.SubElement(trade, "TradeNo").text = "1"
    ET.SubElement(trade, "TradeName").text = get_trade_name()
    areas = ET.SubElement(trade, "Areas")

    processed = 0
    for ws in wb.worksheets:
        title = normalise_sheet_name(ws.title)
        if not title or title in IGNORE_SHEETS:
            continue

        area_no = len(areas.findall("Area")) + 1
        area = ET.SubElement(areas, "Area")
        ET.SubElement(area, "AreaNo").text = str(area_no)
        ET.SubElement(area, "AreaName").text = title
        ET.SubElement(area, "Annotation").text = f"Aus Excel-Tabelle: {title}"

        processes = ET.SubElement(area, "Processes")
        process_no = 1
        process = ET.SubElement(processes, "Process")
        ET.SubElement(process, "ProcessId").text = str(process_no)
        ET.SubElement(process, "ProcessName").text = f"Gefährdungsfälle {title}"
        ET.SubElement(process, "Annotation").text = "Aus dem Excel-Arbeitsblatt übernommen"
        hazards = ET.SubElement(process, "Hazards")

        hazard_count = 0
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
            if len(row) < 10:
                continue
            group = clean_text(row[0])
            lfd = row[1]
            nr = row[2]
            hazard_text = clean_text(row[3])
            risk = clean_text(row[4])
            measure_text = clean_text(row[5])
            action_needed = extract_action_needed(row[7] if len(row) > 7 else None)
            due_source = row[9] if len(row) > 9 else None
            due = parse_date(due_source)
            responsible = extract_responsible(due_source)
            if not hazard_text and not measure_text and not risk:
                continue
            if isinstance(lfd, (int, float)) and (hazard_text or measure_text or risk):
                hazard_count += 1
                hazard = ET.SubElement(hazards, "Hazard")
                ET.SubElement(hazard, "HazardNumber").text = str(int(lfd))
                ET.SubElement(hazard, "HazardName").text = hazard_text or group or f"Gefährdung {int(lfd)}"
                ET.SubElement(hazard, "Risk").text = risk or "mittel"

                measures = ET.SubElement(hazard, "Measures")
                measure = ET.SubElement(measures, "Measure")
                ET.SubElement(measure, "MeasureText").text = measure_text or hazard_text or "Maßnahme gemäß Gefährdungsbeurteilung"
                ET.SubElement(measure, "ActionNeeded").text = str(action_needed).lower()
                ET.SubElement(measure, "DueDate").text = due
                ET.SubElement(measure, "Responsible").text = responsible
                ET.SubElement(measure, "ActualDate").text = ""
                ET.SubElement(measure, "Confirmation").text = ""
                processed += 1

        if hazard_count == 0:
            # remove empty process/area if no hazards were found
            area.remove(processes)
            areas.remove(area)

    ET.indent(root, space="    ")
    tree = ET.ElementTree(root)
    OUTPUT_XML.parent.mkdir(parents=True, exist_ok=True)
    tree.write(OUTPUT_XML, encoding="utf-8", xml_declaration=True)
    print(f"Generated {OUTPUT_XML} with {processed} hazards")


if __name__ == "__main__":
    main()
